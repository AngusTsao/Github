import io
import os
import fitz  # PyMuPDF
from flask import Flask, render_template, request, send_file, jsonify
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB limit

POINTS_TO_EMU = 914400 / 72  # 12700 EMU per point


def pdf_to_pptx(pdf_bytes: bytes) -> bytes:
    """Convert PDF to PPTX with text and image layer separation.

    For each page:
    1. Redact text on a copy to produce a background-only image (image layer).
    2. Extract text blocks with positions from the original page (text layer).
    3. In the PPTX slide, place the background image first, then overlay
       transparent text boxes so the text remains fully editable.
    """
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    prs = Presentation()

    if len(doc) == 0:
        doc.close()
        return b""

    # Set presentation dimensions from first page
    first_rect = doc[0].rect
    prs.slide_width = Emu(int(first_rect.width * POINTS_TO_EMU))
    prs.slide_height = Emu(int(first_rect.height * POINTS_TO_EMU))

    blank_layout = prs.slide_layouts[6]  # blank slide layout

    for page_num in range(len(doc)):
        page = doc[page_num]

        # --- Image layer: render page without text ---
        temp_doc = fitz.open()
        temp_doc.insert_pdf(doc, from_page=page_num, to_page=page_num)
        temp_page = temp_doc[0]

        text_dict = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)

        # Redact all text blocks so the background image has no text
        for block in text_dict.get("blocks", []):
            if block.get("type") == 0:  # text block
                temp_page.add_redact_annot(
                    fitz.Rect(block["bbox"]),
                    fill=(1, 1, 1),  # white fill to erase text
                )
        temp_page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)

        zoom = 2  # 2× zoom for sharper background
        pix = temp_page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        bg_bytes = pix.tobytes("png")
        temp_doc.close()

        # --- Build PPTX slide ---
        slide = prs.slides.add_slide(blank_layout)
        slide_w = prs.slide_width
        slide_h = prs.slide_height

        # Place background image (image layer)
        slide.shapes.add_picture(io.BytesIO(bg_bytes), 0, 0, slide_w, slide_h)

        # Overlay editable text boxes (text layer)
        for block in text_dict.get("blocks", []):
            if block.get("type") != 0:
                continue

            bx0, by0, bx1, by1 = block["bbox"]
            left = Emu(int(bx0 * POINTS_TO_EMU))
            top = Emu(int(by0 * POINTS_TO_EMU))
            width = Emu(max(1, int((bx1 - bx0) * POINTS_TO_EMU)))
            height = Emu(max(1, int((by1 - by0) * POINTS_TO_EMU)))

            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.word_wrap = False

            first_para = True
            for line in block.get("lines", []):
                if first_para:
                    p = tf.paragraphs[0]
                    first_para = False
                else:
                    p = tf.add_paragraph()

                for span in line.get("spans", []):
                    run = p.add_run()
                    run.text = span.get("text", "")

                    font = run.font
                    font.size = Pt(span.get("size", 12))

                    # Decode packed RGB integer (PyMuPDF stores color as 0xRRGGBB
                    # where RR=red byte, GG=green byte, BB=blue byte)
                    color_int = span.get("color", 0)
                    r = (color_int >> 16) & 0xFF
                    g = (color_int >> 8) & 0xFF
                    b = color_int & 0xFF
                    font.color.rgb = RGBColor(r, g, b)

                    # Font style flags: position 4 (1 << 4 = 16) = bold,
                    # position 1 (1 << 1 = 2) = italic
                    flags = span.get("flags", 0)
                    font.bold = bool(flags & (1 << 4))
                    font.italic = bool(flags & (1 << 1))

                    font.name = span.get("font", "Calibri") or "Calibri"

    doc.close()

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "未選擇檔案"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "未選擇檔案"}), 400

    if not file.filename.lower().endswith(".pdf"):
        return jsonify({"error": "僅支援 PDF 檔案"}), 400

    pdf_bytes = file.read()
    if len(pdf_bytes) == 0:
        return jsonify({"error": "檔案為空"}), 400

    pptx_bytes = pdf_to_pptx(pdf_bytes)
    if not pptx_bytes:
        return jsonify({"error": "PDF 轉換失敗，請確認檔案是否為有效的 PDF"}), 500

    base_name = os.path.splitext(file.filename)[0]
    output_filename = f"{base_name}.pptx"

    return send_file(
        io.BytesIO(pptx_bytes),
        as_attachment=True,
        download_name=output_filename,
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


if __name__ == "__main__":
    app.run(debug=False)
